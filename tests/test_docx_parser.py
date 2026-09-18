"""
Unit tests for DocxTranslator (Microsoft Word .docx).
"""

import tempfile
import unittest
from pathlib import Path
from docx import Document

from config import AppConfig
from docx_parser import DocxTranslator
from translation_cache import TranslationCache


class MockLLMManager:
    def __init__(self):
        self.call_count = 0

    def translate_batch(self, texts, target_lang="Vietnamese", *args, **kwargs):
        self.call_count += 1
        return [f"[{target_lang}] {t}" for t in texts]


class TestDocxTranslator(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.test_docx_path = Path(self.temp_dir.name) / "sample.docx"

        # Create a real docx document with paragraphs and tables
        doc = Document()
        doc.add_heading("User Guide", level=1)
        doc.add_paragraph("Welcome to the application user guide.")
        doc.add_paragraph("Please read all instructions carefully.")

        # Add table
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Feature"
        table.cell(0, 1).text = "Status"
        table.cell(1, 0).text = "Login Authentication"
        table.cell(1, 1).text = "Ready"

        doc.save(str(self.test_docx_path))

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_inspect_file(self):
        info = DocxTranslator.inspect_file(str(self.test_docx_path))
        self.assertEqual(info["type"], "docx")
        self.assertGreater(info["paragraphs"], 0)
        self.assertEqual(info["tables"], 1)
        self.assertGreater(info["word_count"], 5)
        self.assertGreater(info["translatable_segments"], 3)

    def test_translation_flow(self):
        config = AppConfig()
        mock_llm = MockLLMManager()
        cache = TranslationCache(db_path=Path(self.temp_dir.name) / "cache.db")
        translator = DocxTranslator(config, mock_llm, cache=cache)

        out_file = translator.process_file(
            file_path=str(self.test_docx_path),
            target_lang="Vietnamese"
        )

        out_path = Path(out_file)
        self.assertTrue(out_path.exists())
        self.assertEqual(out_path.name, "sample_vi.docx")

        # Verify translated content in docx
        translated_doc = Document(str(out_path))
        all_texts = [p.text for p in translated_doc.paragraphs]
        for t in translated_doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        all_texts.append(p.text)

        joined = " ".join(all_texts)
        self.assertIn("[Vietnamese] User Guide", joined)
        self.assertIn("[Vietnamese] Welcome to the application user guide.", joined)
        self.assertIn("[Vietnamese] Login Authentication", joined)

    def test_caching(self):
        config = AppConfig()
        mock_llm = MockLLMManager()
        cache = TranslationCache(db_path=Path(self.temp_dir.name) / "cache2.db")
        translator = DocxTranslator(config, mock_llm, cache=cache)

        # 1st run: fills cache
        translator.process_file(file_path=str(self.test_docx_path), target_lang="Vietnamese")
        first_calls = mock_llm.call_count
        self.assertGreater(first_calls, 0)

        # 2nd run: should hit cache, call_count should not increase
        translator.process_file(file_path=str(self.test_docx_path), target_lang="Vietnamese")
        self.assertEqual(mock_llm.call_count, first_calls)


if __name__ == "__main__":
    unittest.main()

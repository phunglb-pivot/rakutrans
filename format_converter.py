"""
Format conversion engine for RakuTrans AI.
Enables cross-format document and data conversions:
- Excel (.xlsx) <-> Markdown (.md)
- Excel (.xlsx) <-> CSV (.csv)
- CSV/TSV (.csv, .tsv) <-> Excel (.xlsx)
- CSV/TSV (.csv, .tsv) <-> Markdown (.md)
- Word (.docx) <-> Markdown (.md)
- Word (.docx) -> Text (.txt)
- PowerPoint (.pptx) -> Markdown (.md)
- PowerPoint (.pptx) -> Text (.txt)
- Markdown (.md) -> Word (.docx)
- Text (.txt) -> Markdown (.md) / Word (.docx)
"""

import csv
import os
import re
from pathlib import Path
from typing import List, Optional, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import get_smart_output_path, get_unique_output_path


# -----------------------------------------------------------------------------
# Excel <-> Markdown
# -----------------------------------------------------------------------------

def excel_to_markdown(xlsx_path: str, md_output_path: Optional[str] = None) -> str:
    """
    Converts an Excel (.xlsx) workbook into a clean, multi-sheet Markdown document.
    Each sheet is represented with a heading and a GitHub-flavored Markdown table.
    """
    src = Path(xlsx_path)
    if not src.exists():
        raise FileNotFoundError(f"Excel file not found: {xlsx_path}")

    if not md_output_path:
        md_output_path = str(get_unique_output_path(src.with_suffix(".md")))

    wb = openpyxl.load_workbook(str(src), data_only=False)
    md_lines: List[str] = []

    doc_title = re.sub(r"^\[Translated\]_", "", src.stem)
    md_lines.append(f"# {doc_title}\n")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        md_lines.append(f"## {sheet_name}\n")

        raw_rows = list(ws.iter_rows(values_only=True))
        filtered_rows = []
        for row in raw_rows:
            if any(c is not None and str(c).strip() != "" for c in row):
                filtered_rows.append(row)

        if not filtered_rows:
            md_lines.append("*(Trống / Empty sheet)*\n")
            continue

        max_cols = max(len(row) for row in filtered_rows)
        active_cols = []
        for col_idx in range(max_cols):
            has_val = any(
                col_idx < len(r) and r[col_idx] is not None and str(r[col_idx]).strip() != ""
                for r in filtered_rows
            )
            if has_val:
                active_cols.append(col_idx)

        if not active_cols:
            md_lines.append("*(Trống / Empty sheet)*\n")
            continue

        def sanitize_cell(val: any) -> str:
            if val is None:
                return ""
            s = str(val).strip()
            s = s.replace("\r\n", "<br>").replace("\n", "<br>")
            s = s.replace("|", "\\|")
            return s

        table_rows: List[List[str]] = []
        for r in filtered_rows:
            row_cells = [sanitize_cell(r[c] if c < len(r) else "") for c in active_cols]
            table_rows.append(row_cells)

        if table_rows:
            header_cells = table_rows[0]
            if all(not c for c in header_cells):
                header_cells = [f"Col {i+1}" for i in range(len(header_cells))]

            md_lines.append("| " + " | ".join(header_cells) + " |")
            md_lines.append("| " + " | ".join(["---"] * len(header_cells)) + " |")

            for data_row in table_rows[1:]:
                md_lines.append("| " + " | ".join(data_row) + " |")

            md_lines.append("")

    out_file = Path(md_output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    with open(out_file, "w", encoding="utf-8") as f:
        f.write("\n".join(md_lines).strip() + "\n")

    return str(out_file)


def markdown_to_excel(md_path: str, xlsx_output_path: Optional[str] = None) -> str:
    """
    Converts a Markdown (.md) document into a styled Excel (.xlsx) workbook.
    """
    src = Path(md_path)
    if not src.exists():
        raise FileNotFoundError(f"Markdown file not found: {md_path}")

    if not xlsx_output_path:
        xlsx_output_path = str(get_unique_output_path(src.with_suffix(".xlsx")))

    with open(src, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()

    wb = openpyxl.Workbook()
    default_sheet = wb.active
    default_sheet.title = "Document"

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    title_font = Font(name="Segoe UI", size=16, bold=True, color="0F243E")
    normal_font = Font(name="Segoe UI", size=10)
    
    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    lines = content.splitlines()
    current_sheet = default_sheet
    current_row = 1
    sheet_created_count = 0

    idx = 0
    while idx < len(lines):
        line = lines[idx].strip()
        if not line:
            idx += 1
            continue

        if line.startswith("# ") and not line.startswith("## "):
            title_text = line[2:].strip()
            cell = current_sheet.cell(row=current_row, column=1, value=title_text)
            cell.font = title_font
            current_row += 2
            idx += 1
            continue

        if line.startswith("## "):
            section_title = line[3:].strip()
            clean_title = re.sub(r'[\\/*?:\[\]]', '_', section_title)[:31]
            if current_row == 1 and sheet_created_count == 0:
                current_sheet.title = clean_title
                sheet_created_count += 1
            else:
                sheet_name = clean_title
                counter = 2
                while sheet_name in wb.sheetnames:
                    sheet_name = f"{clean_title[:28]}_{counter}"
                    counter += 1
                current_sheet = wb.create_sheet(title=sheet_name)
                current_row = 1
                sheet_created_count += 1

            cell = current_sheet.cell(row=current_row, column=1, value=section_title)
            cell.font = Font(name="Segoe UI", size=14, bold=True, color="1F4E79")
            current_row += 2
            idx += 1
            continue

        if line.startswith("|") and line.endswith("|"):
            table_lines = []
            while idx < len(lines) and lines[idx].strip().startswith("|") and lines[idx].strip().endswith("|"):
                table_lines.append(lines[idx].strip())
                idx += 1

            table_data = []
            for t_line in table_lines:
                inner = t_line.strip("|").strip()
                if set(inner.replace("|", "").replace("-", "").replace(":", "").strip()) == set():
                    continue

                cells = [c.strip().replace("<br>", "\n").replace("\\|", "|") for c in t_line.strip("|").split("|")]
                table_data.append(cells)

            if table_data:
                for r_idx, row_values in enumerate(table_data):
                    is_header = (r_idx == 0)
                    for c_idx, val in enumerate(row_values):
                        cell = current_sheet.cell(row=current_row, column=c_idx + 1)
                        if val.isdigit():
                            cell.value = int(val)
                        else:
                            try:
                                cell.value = float(val) if "." in val else val
                            except ValueError:
                                cell.value = val

                        cell.border = thin_border
                        if is_header:
                            cell.font = header_font
                            cell.fill = header_fill
                            cell.alignment = center_align
                        else:
                            cell.font = normal_font
                            cell.alignment = left_align

                    current_row += 1

                current_row += 2
            continue

        cell = current_sheet.cell(row=current_row, column=1, value=line)
        cell.font = normal_font
        cell.alignment = left_align
        current_row += 1
        idx += 1

    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val:
                    lines_in_cell = str(val).split("\n")
                    longest = max(len(l) for l in lines_in_cell)
                    if longest > max_len:
                        max_len = longest
            ws.column_dimensions[col_letter].width = max(12, min(max_len + 4, 60))

    out_file = Path(xlsx_output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_file))
    return str(out_file)


# -----------------------------------------------------------------------------
# CSV / TSV Conversions
# -----------------------------------------------------------------------------

def _detect_csv_params(file_path: Path) -> Tuple[str, str]:
    """Detects encoding and delimiter for CSV/TSV."""
    enc = "utf-8"
    for candidate_enc in ["utf-8-sig", "utf-8", "shift_jis", "cp1252", "latin-1"]:
        try:
            with open(file_path, "r", encoding=candidate_enc) as f:
                f.read(4096)
                enc = candidate_enc
                break
        except (UnicodeDecodeError, LookupError):
            continue

    delim = "\t" if file_path.suffix.lower() == ".tsv" else ","
    try:
        with open(file_path, "r", encoding=enc, errors="replace") as f:
            sample = f.read(4096)
            dialect = csv.Sniffer().sniff(sample, delimiters=[",", "\t", ";", "|"])
            delim = dialect.delimiter
    except Exception:
        pass

    return enc, delim


def csv_to_excel(csv_path: str, xlsx_output_path: Optional[str] = None) -> str:
    """Converts a CSV or TSV file into a formatted Excel (.xlsx) workbook."""
    src = Path(csv_path)
    if not src.exists():
        raise FileNotFoundError(f"CSV file not found: {csv_path}")

    if not xlsx_output_path:
        xlsx_output_path = str(get_unique_output_path(src.with_suffix(".xlsx")))

    enc, delim = _detect_csv_params(src)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = src.stem[:31]

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    normal_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    center_align = Alignment(horizontal="center", vertical="center")

    with open(src, "r", encoding=enc, errors="replace") as f:
        reader = csv.reader(f, delimiter=delim)
        for r_idx, row in enumerate(reader, start=1):
            is_header = (r_idx == 1)
            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx)
                val_str = str(val).strip()
                if not is_header and val_str.isdigit():
                    cell.value = int(val_str)
                elif not is_header:
                    try:
                        cell.value = float(val_str) if "." in val_str else val
                    except ValueError:
                        cell.value = val
                else:
                    cell.value = val

                cell.border = thin_border
                if is_header:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center_align
                else:
                    cell.font = normal_font
                    cell.alignment = left_align

    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col) if col else 10
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(12, min(max_len + 4, 60))

    out_file = Path(xlsx_output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_file))
    return str(out_file)


def excel_to_csv(xlsx_path: str, csv_output_path: Optional[str] = None) -> str:
    """Exports active sheet of an Excel (.xlsx) workbook to CSV with UTF-8-SIG."""
    src = Path(xlsx_path)
    if not src.exists():
        raise FileNotFoundError(f"Excel file not found: {xlsx_path}")

    if not csv_output_path:
        csv_output_path = str(get_unique_output_path(src.with_suffix(".csv")))

    wb = openpyxl.load_workbook(str(src), data_only=True)
    ws = wb.active

    out_file = Path(csv_output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)

    with open(out_file, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            if any(c is not None and str(c).strip() != "" for c in row):
                writer.writerow([c if c is not None else "" for c in row])

    return str(out_file)


def csv_to_markdown(csv_path: str, md_output_path: Optional[str] = None) -> str:
    """Converts CSV or TSV file into a GitHub Flavored Markdown table."""
    src = Path(csv_path)
    if not src.exists():
        raise FileNotFoundError(f"CSV file not found: {csv_path}")

    if not md_output_path:
        md_output_path = str(get_unique_output_path(src.with_suffix(".md")))

    enc, delim = _detect_csv_params(src)

    rows = []
    with open(src, "r", encoding=enc, errors="replace") as f:
        reader = csv.reader(f, delimiter=delim)
        for r in reader:
            if any(c.strip() for c in r):
                rows.append([c.strip().replace("\r\n", "<br>").replace("\n", "<br>").replace("|", "\\|") for c in r])

    md_lines = [f"# {src.stem}\n"]
    if rows:
        header = rows[0]
        md_lines.append("| " + " | ".join(header) + " |")
        md_lines.append("| " + " | ".join(["---"] * len(header)) + " |")
        for data_row in rows[1:]:
            padded = data_row + [""] * max(0, len(header) - len(data_row))
            md_lines.append("| " + " | ".join(padded[:len(header)]) + " |")
    else:
        md_lines.append("*(Empty CSV)*")

    out_file = Path(md_output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    with open(out_file, "w", encoding="utf-8") as f:
        f.write("\n".join(md_lines) + "\n")

    return str(out_file)


def markdown_to_csv(md_path: str, csv_output_path: Optional[str] = None) -> str:
    """Extracts tables from a Markdown file and writes to CSV."""
    src = Path(md_path)
    if not src.exists():
        raise FileNotFoundError(f"Markdown file not found: {md_path}")

    if not csv_output_path:
        csv_output_path = str(get_unique_output_path(src.with_suffix(".csv")))

    with open(src, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()

    lines = content.splitlines()
    table_rows = []
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            inner = stripped.strip("|").strip()
            if set(inner.replace("|", "").replace("-", "").replace(":", "").strip()) == set():
                continue
            cells = [c.strip().replace("<br>", "\n").replace("\\|", "|") for c in stripped.strip("|").split("|")]
            table_rows.append(cells)

    out_file = Path(csv_output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    with open(out_file, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(table_rows)

    return str(out_file)


# -----------------------------------------------------------------------------
# Word (.docx) <-> Markdown & Text
# -----------------------------------------------------------------------------

def docx_to_markdown(docx_path: str, md_output_path: Optional[str] = None) -> str:
    """Converts a Microsoft Word (.docx) document to GitHub Flavored Markdown."""
    src = Path(docx_path)
    if not src.exists():
        raise FileNotFoundError(f"Word document not found: {docx_path}")

    if not md_output_path:
        md_output_path = str(get_unique_output_path(src.with_suffix(".md")))

    import docx
    doc = docx.Document(str(src))
    md_lines: List[str] = [f"# {src.stem}\n"]

    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        style_name = p.style.name.lower() if p.style else ""
        if "heading 1" in style_name:
            md_lines.append(f"# {txt}\n")
        elif "heading 2" in style_name:
            md_lines.append(f"## {txt}\n")
        elif "heading 3" in style_name:
            md_lines.append(f"### {txt}\n")
        elif "bullet" in style_name or "list" in style_name:
            md_lines.append(f"- {txt}")
        else:
            md_lines.append(f"{txt}\n")

    for t in doc.tables:
        t_rows = []
        for r in t.rows:
            t_rows.append([cell.text.strip().replace("\n", "<br>").replace("|", "\\|") for cell in r.cells])
        if t_rows:
            header = t_rows[0]
            md_lines.append("| " + " | ".join(header) + " |")
            md_lines.append("| " + " | ".join(["---"] * len(header)) + " |")
            for row in t_rows[1:]:
                md_lines.append("| " + " | ".join(row) + " |")
            md_lines.append("")

    out_file = Path(md_output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    with open(out_file, "w", encoding="utf-8") as f:
        f.write("\n".join(md_lines).strip() + "\n")

    return str(out_file)


def markdown_to_docx(md_path: str, docx_output_path: Optional[str] = None) -> str:
    """Converts a Markdown (.md) document into a styled Word (.docx) document."""
    src = Path(md_path)
    if not src.exists():
        raise FileNotFoundError(f"Markdown file not found: {md_path}")

    if not docx_output_path:
        docx_output_path = str(get_unique_output_path(src.with_suffix(".docx")))

    import docx
    doc = docx.Document()

    with open(src, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()

    lines = content.splitlines()
    idx = 0
    while idx < len(lines):
        line = lines[idx].rstrip()
        stripped = line.strip()
        if not stripped:
            idx += 1
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
            idx += 1
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
            idx += 1
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
            idx += 1
        elif stripped.startswith("- ") or stripped.startswith("* "):
            doc.add_paragraph(stripped[2:], style='List Bullet')
            idx += 1
        elif stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while idx < len(lines) and lines[idx].strip().startswith("|") and lines[idx].strip().endswith("|"):
                table_lines.append(lines[idx].strip())
                idx += 1

            parsed_rows = []
            for tl in table_lines:
                inner = tl.strip("|").strip()
                if set(inner.replace("|", "").replace("-", "").replace(":", "").strip()) == set():
                    continue
                cells = [c.strip().replace("<br>", "\n").replace("\\|", "|") for c in tl.strip("|").split("|")]
                parsed_rows.append(cells)

            if parsed_rows:
                num_cols = max(len(r) for r in parsed_rows)
                t = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                t.style = 'Table Grid'
                for r_idx, r_data in enumerate(parsed_rows):
                    for c_idx, val in enumerate(r_data):
                        cell = t.cell(r_idx, c_idx)
                        cell.text = val
                        if r_idx == 0:
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    run.bold = True
        elif stripped.startswith("```"):
            code_lines = []
            idx += 1
            while idx < len(lines) and not lines[idx].strip().startswith("```"):
                code_lines.append(lines[idx])
                idx += 1
            if idx < len(lines) and lines[idx].strip().startswith("```"):
                idx += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Courier New"
        else:
            doc.add_paragraph(stripped)
            idx += 1

    out_file = Path(docx_output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_file))
    return str(out_file)


def docx_to_text(docx_path: str, txt_output_path: Optional[str] = None) -> str:
    """Extracts text from Word document into plain text (.txt)."""
    src = Path(docx_path)
    if not src.exists():
        raise FileNotFoundError(f"Word document not found: {docx_path}")

    if not txt_output_path:
        txt_output_path = str(get_unique_output_path(src.with_suffix(".txt")))

    import docx
    doc = docx.Document(str(src))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        for row in t.rows:
            row_text = "\t".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                texts.append(row_text)

    out_file = Path(txt_output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    with open(out_file, "w", encoding="utf-8") as f:
        f.write("\n\n".join(texts) + "\n")

    return str(out_file)


# -----------------------------------------------------------------------------
# PowerPoint (.pptx) -> Markdown & Text
# -----------------------------------------------------------------------------

def pptx_to_markdown(pptx_path: str, md_output_path: Optional[str] = None) -> str:
    """Extracts slides outline and tables from PowerPoint (.pptx) to Markdown."""
    src = Path(pptx_path)
    if not src.exists():
        raise FileNotFoundError(f"PowerPoint presentation not found: {pptx_path}")

    if not md_output_path:
        md_output_path = str(get_unique_output_path(src.with_suffix(".md")))

    from pptx import Presentation
    prs = Presentation(str(src))
    md_lines: List[str] = [f"# {src.stem}\n"]

    for idx, slide in enumerate(prs.slides, start=1):
        md_lines.append(f"## Slide {idx}\n")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    txt = p.text.strip()
                    if txt:
                        md_lines.append(f"- {txt}")
                md_lines.append("")
            elif shape.has_table:
                t_rows = []
                for row in shape.table.rows:
                    t_rows.append([c.text.strip().replace("\n", "<br>").replace("|", "\\|") for c in row.cells])
                if t_rows:
                    header = t_rows[0]
                    md_lines.append("| " + " | ".join(header) + " |")
                    md_lines.append("| " + " | ".join(["---"] * len(header)) + " |")
                    for row in t_rows[1:]:
                        md_lines.append("| " + " | ".join(row) + " |")
                    md_lines.append("")

    out_file = Path(md_output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    with open(out_file, "w", encoding="utf-8") as f:
        f.write("\n".join(md_lines).strip() + "\n")

    return str(out_file)


def pptx_to_text(pptx_path: str, txt_output_path: Optional[str] = None) -> str:
    """Extracts text from PowerPoint (.pptx) into plain text."""
    src = Path(pptx_path)
    if not src.exists():
        raise FileNotFoundError(f"PowerPoint presentation not found: {pptx_path}")

    if not txt_output_path:
        txt_output_path = str(get_unique_output_path(src.with_suffix(".txt")))

    from pptx import Presentation
    prs = Presentation(str(src))
    texts: List[str] = []

    for idx, slide in enumerate(prs.slides, start=1):
        texts.append(f"=== Slide {idx} ===")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if p.text.strip():
                        texts.append(p.text.strip())
            elif shape.has_table:
                for row in shape.table.rows:
                    row_txt = "\t".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_txt:
                        texts.append(row_txt)
        texts.append("")

    out_file = Path(txt_output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    with open(out_file, "w", encoding="utf-8") as f:
        f.write("\n".join(texts) + "\n")

    return str(out_file)


# -----------------------------------------------------------------------------
# Plain Text Conversions
# -----------------------------------------------------------------------------

def text_to_markdown(txt_path: str, md_output_path: Optional[str] = None) -> str:
    """Converts a Plain Text (.txt) file into Markdown (.md)."""
    src = Path(txt_path)
    if not src.exists():
        raise FileNotFoundError(f"Text file not found: {txt_path}")

    if not md_output_path:
        md_output_path = str(get_unique_output_path(src.with_suffix(".md")))

    with open(src, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()

    out_file = Path(md_output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    with open(out_file, "w", encoding="utf-8") as f:
        f.write(f"# {src.stem}\n\n" + content.strip() + "\n")

    return str(out_file)


# -----------------------------------------------------------------------------
# Dispatcher & Option Registry
# -----------------------------------------------------------------------------

def get_convert_options(file_ext: str) -> List[str]:
    """
    Returns available target format keys for a given source file extension.
    E.g.:
      '.xlsx' -> ['md', 'csv']
      '.csv', '.tsv' -> ['xlsx', 'md']
      '.docx' -> ['md', 'txt']
      '.pptx' -> ['md', 'txt']
      '.md' -> ['xlsx', 'docx', 'csv']
      '.txt' -> ['md', 'docx']
    """
    ext = file_ext.lower().lstrip(".")
    if ext == "xlsx":
        return ["md", "csv"]
    elif ext in ["csv", "tsv"]:
        return ["xlsx", "md"]
    elif ext == "docx":
        return ["md", "txt"]
    elif ext == "pptx":
        return ["md", "txt"]
    elif ext == "md":
        return ["xlsx", "docx", "csv"]
    elif ext == "txt":
        return ["md", "docx"]
    return []


def convert_file_direct(source_path: str, target_format: str) -> str:
    """
    Directly converts a document to target_format without translation.
    target_format: 'xlsx', 'md', 'docx', 'csv', 'txt'
    Returns the path to the converted output file.
    """
    path = Path(source_path)
    src_fmt = path.suffix.lower().lstrip(".")
    tgt_fmt = target_format.lower().strip(".")

    out_path = str(get_smart_output_path(source_path=path, target_format=f".{tgt_fmt}"))

    key = (src_fmt, tgt_fmt)
    if key == ("xlsx", "md"):
        return excel_to_markdown(str(path), out_path)
    elif key == ("xlsx", "csv"):
        return excel_to_csv(str(path), out_path)
    elif key in [("csv", "xlsx"), ("tsv", "xlsx")]:
        return csv_to_excel(str(path), out_path)
    elif key in [("csv", "md"), ("tsv", "md")]:
        return csv_to_markdown(str(path), out_path)
    elif key == ("md", "xlsx"):
        return markdown_to_excel(str(path), out_path)
    elif key == ("md", "docx"):
        return markdown_to_docx(str(path), out_path)
    elif key == ("md", "csv"):
        return markdown_to_csv(str(path), out_path)
    elif key == ("docx", "md"):
        return docx_to_markdown(str(path), out_path)
    elif key == ("docx", "txt"):
        return docx_to_text(str(path), out_path)
    elif key == ("pptx", "md"):
        return pptx_to_markdown(str(path), out_path)
    elif key == ("pptx", "txt"):
        return pptx_to_text(str(path), out_path)
    elif key == ("txt", "md"):
        return text_to_markdown(str(path), out_path)
    elif key == ("txt", "docx"):
        return markdown_to_docx(str(path), out_path)

    raise ValueError(f"Cannot convert .{src_fmt} to .{tgt_fmt}")


def convert_translated_output(
    translated_file: str,
    target_format: str = "auto"
) -> str:
    """
    Checks if post-translation conversion is requested:
    - 'auto': leaves file in original format.
    - Other target format: converts if conversion is supported for the file extension.
    Returns the path to the final output file.
    """
    path = Path(translated_file)
    fmt = (target_format or "auto").lower().lstrip(".")

    if fmt == "auto" or not fmt:
        return str(path)

    src_ext = path.suffix.lower().lstrip(".")
    if fmt == src_ext:
        return str(path)

    valid_targets = get_convert_options(src_ext)
    if fmt in valid_targets:
        return convert_file_direct(str(path), fmt)

    return str(path)

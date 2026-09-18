"""
Excel parsing and translation engine for RakuTrans AI.
Uses openpyxl to translate multi-sheet workbooks while strictly preserving
all formulas, formatting, styles, dimensions, and structural properties.
"""

import os
from pathlib import Path
from typing import List, Tuple, Dict, Callable, Optional
from concurrent.futures import ThreadPoolExecutor, as_completed
import openpyxl

import re
from config import AppConfig, TranslationOptions, get_smart_output_path
from llm_manager import LLMManager, LLMError
from translation_cache import TranslationCache


from base_translator import BaseTranslator, safe_save_file


def safe_save_workbook(workbook: openpyxl.Workbook, target_path: Path) -> Path:
    """
    Saves workbook to target_path. If target_path is locked (e.g. open in Microsoft Excel),
    automatically falls back to target_path (1).xlsx, target_path (2).xlsx to avoid crashing.
    """
    return safe_save_file(lambda p: workbook.save(str(p)), target_path)


class ExcelTranslator(BaseTranslator):
    """
    Handles translation of Excel (.xlsx) files.
    Preserves formulas, sheet order, cell styles, merged cells, and formatting.
    Features in-file deduplication and local Translation Memory to avoid redundant API calls.
    Supports concurrent multi-threaded batching for maximum translation speed.
    Supports safe-save against Excel file locking and optional bilingual sheet export.
    """

    def __init__(
        self,
        config: AppConfig,
        llm_manager: LLMManager,
        cache: Optional[TranslationCache] = None
    ):
        super().__init__(config, llm_manager, cache)

    def cancel(self) -> None:
        """Signals the translation process to abort."""
        self._is_cancelled = True

    @staticmethod
    def should_translate_cell(value: any) -> bool:
        """
        Determines whether a cell value should be translated.
        Strictly ignores numbers, booleans, empty cells, and formulas.
        """
        if value is None:
            return False

        # Strictly ignore non-string types (int, float, datetime, bool, etc.)
        if not isinstance(value, str):
            return False

        trimmed = value.strip()
        if not trimmed:
            return False

        # Strictly ignore formula cells starting with '='
        if trimmed.startswith("="):
            return False

        # Ignore strings that are purely numeric (e.g. "12345", "0.95")
        try:
            float(trimmed.replace(",", ""))
            return False
        except ValueError:
            pass

        return True

    @classmethod
    def inspect_file(cls, file_path: str) -> dict:
        """
        Quickly inspects an Excel workbook to extract metadata:
        sheet count, translatable cell count, and formula count.
        """
        path = Path(file_path)
        if not path.exists():
            return {"error": "File not found"}

        try:
            wb = openpyxl.load_workbook(filename=file_path, data_only=False, read_only=False)
            sheet_names = wb.sheetnames
            translatable_cells = 0
            formula_count = 0
            unique_set = set()

            for name in sheet_names:
                ws = wb[name]
                for row in ws.iter_rows():
                    for cell in row:
                        val = cell.value
                        if val is not None and isinstance(val, str) and val.strip().startswith("="):
                            formula_count += 1
                        elif cls.should_translate_cell(val):
                            translatable_cells += 1
                            unique_set.add(val)

            wb.close()
            return {
                "type": "excel",
                "sheet_count": len(sheet_names),
                "sheet_names": sheet_names,
                "translatable_cells": translatable_cells,
                "unique_texts": len(unique_set),
                "formula_count": formula_count
            }
        except Exception as e:
            return {"type": "excel", "error": str(e)}

    def process_file(
        self,
        file_path: str,
        target_lang: Optional[str] = None,
        source_lang: Optional[str] = "auto",
        options: Optional[TranslationOptions] = None,
        progress_callback: Optional[Callable[[int, int, str], None]] = None,
        custom_output_path: Optional[str] = None,
        *args,
        **kwargs
    ) -> str:
        """
        Main processing method for Excel files.
        Extracts translatable text cells, deduplicates unique strings,
        batches uncached items, translates via LLM (with concurrent batching),
        and re-injects translations while strictly preserving formulas and styling.
        Supports safe saving against file lock and optional bilingual sheet duplication.
        Returns the output file path.
        """
        # Resolve target_lang and source_lang flexibly
        if target_lang is None:
            target_lang = kwargs.get("target_lang", "Vietnamese")

        if isinstance(source_lang, TranslationOptions):
            options = source_lang
            source_lang = "auto"
        elif isinstance(source_lang, str) and source_lang != "auto" and "target_lang" not in kwargs and len(args) > 0:
            # Legacy positional: (file_path, source_lang, target_lang, options, ...)
            legacy_src = target_lang
            legacy_tgt = source_lang
            target_lang = legacy_tgt
            source_lang = legacy_src
            if args:
                options = args[0]
        elif source_lang is None:
            source_lang = "auto"

        self._is_cancelled = False
        src_path = Path(file_path)
        if not src_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        # Strict requirement: data_only=False preserves formulas (=SUM, etc.)
        if progress_callback:
            progress_callback(0, 100, "Loading workbook structure & styles...")

        workbook = openpyxl.load_workbook(filename=file_path, data_only=False)

        # 1. Scan all sheets and extract translatable cells
        # Format: (sheet_name, cell_coordinate, original_text)
        translatable_cells: List[Tuple[str, str, str]] = []

        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if self.should_translate_cell(cell.value):
                        translatable_cells.append((sheet_name, cell.coordinate, cell.value))

        total_cells = len(translatable_cells)
        if total_cells == 0:
            target_path = Path(custom_output_path) if custom_output_path else get_smart_output_path(
                source_path=src_path,
                target_lang=target_lang,
                is_bilingual=bool(options and getattr(options, "bilingual_mode", False))
            )
            out_file = safe_save_workbook(workbook, target_path)
            return str(out_file)

        # 2. In-file Deduplication: Extract unique texts preserving original discovery order
        unique_texts: List[str] = list(dict.fromkeys([item[2] for item in translatable_cells]))
        unique_count = len(unique_texts)
        duplicate_saved = total_cells - unique_count

        if progress_callback:
            progress_callback(
                5, 100, 
                f"Deduplication: {unique_count} unique texts ({duplicate_saved} duplicate cells saved)..."
            )

        # 3. Check local translation memory cache for unique texts
        cached_map, uncached_items = self.cache.get_batch(
            texts=unique_texts,
            target_lang=target_lang,
            options=options,
            source_lang=source_lang
        )
        cache_hits = len(cached_map)
        unique_translations: Dict[str, str] = {}

        # Populate cached results
        for idx, translated_text in cached_map.items():
            unique_translations[unique_texts[idx]] = translated_text

        # 4. Batch uncached unique texts
        if uncached_items:
            batch_size = max(10, min(self.config.batch_size, 100))
            batches: List[List[Tuple[int, str]]] = [
                uncached_items[i:i + batch_size]
                for i in range(0, len(uncached_items), batch_size)
            ]
            total_batches = len(batches)
            completed_batches = 0

            # Worker function for single batch
            def translate_single_batch(batch_tuple: Tuple[int, List[Tuple[int, str]]]) -> Tuple[List[str], List[str]]:
                b_idx, b_items = batch_tuple
                if self._is_cancelled:
                    raise InterruptedError("Translation was cancelled by user.")
                
                texts_to_translate = [item[1] for item in b_items]
                translated = self.llm_manager.translate_batch(
                    texts=texts_to_translate,
                    target_lang=target_lang,
                    options=options,
                    source_lang=source_lang
                )
                # Persist to cache
                self.cache.save_batch(
                    source_texts=texts_to_translate,
                    translated_texts=translated,
                    target_lang=target_lang,
                    options=options,
                    source_lang=source_lang
                )
                return texts_to_translate, translated

            # Concurrency: Use ThreadPoolExecutor (up to 3 concurrent batch requests)
            max_workers = min(3, total_batches) if total_batches > 1 else 1

            with ThreadPoolExecutor(max_workers=max_workers) as executor:
                indexed_batches = list(enumerate(batches))
                future_to_batch = {
                    executor.submit(translate_single_batch, ib): ib[0]
                    for ib in indexed_batches
                }

                for future in as_completed(future_to_batch):
                    if self._is_cancelled:
                        executor.shutdown(wait=False, cancel_futures=True)
                        raise InterruptedError("Translation was cancelled by user.")

                    src_texts, trans_texts = future.result()
                    for s, t in zip(src_texts, trans_texts):
                        unique_translations[s] = t

                    completed_batches += 1
                    pct = 10 + int((completed_batches / total_batches) * 80)
                    msg = (
                        f"Batch {completed_batches}/{total_batches} completed "
                        f"({cache_hits} cached | {duplicate_saved} deduplicated)"
                    )
                    if progress_callback:
                        progress_callback(pct, 100, msg)

        if self._is_cancelled:
            raise InterruptedError("Translation was cancelled by user.")

        # 4.5. Optional Bilingual Mode: Duplicate original sheets before in-place translation
        if options and getattr(options, "bilingual_mode", False):
            if progress_callback:
                progress_callback(90, 100, "Bilingual mode: Archiving original worksheets...")
            orig_sheetnames = list(workbook.sheetnames)
            for s_name in orig_sheetnames:
                orig_ws = workbook[s_name]
                copy_ws = workbook.copy_worksheet(orig_ws)
                copy_ws.title = f"{s_name[:24]}_Orig"

        # 5. Re-inject translated text back into workbook cell coordinates
        if progress_callback:
            progress_callback(92, 100, f"Re-injecting {total_cells} cells into workbook...")

        for sheet_name, coord, orig_val in translatable_cells:
            ws = workbook[sheet_name]
            ws[coord].value = unique_translations.get(orig_val, orig_val)

        # 6. Safe Save workbook (auto-increments if file locked by Excel)
        if progress_callback:
            progress_callback(96, 100, "Saving translated workbook...")

        target_path = Path(custom_output_path) if custom_output_path else get_smart_output_path(
            source_path=src_path,
            target_lang=target_lang,
            is_bilingual=bool(options and getattr(options, "bilingual_mode", False))
        )
        output_path = safe_save_workbook(workbook, target_path)

        if progress_callback:
            note = f"Saved to {output_path.name}"
            if output_path.name != target_path.name:
                note += " (safe-renamed to avoid file lock)"
            progress_callback(100, 100, f"Completed! {note}")

        return str(output_path)

"""
Unit tests for PptxTranslator (Microsoft PowerPoint .pptx).
"""

import tempfile
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

from config import AppConfig
from pptx_parser import PptxTranslator
from translation_cache import TranslationCache


class MockLLMManager:
    def __init__(self):
        self.call_count = 0

    def translate_batch(self, texts, target_lang="Vietnamese", *args, **kwargs):
        self.call_count += 1
        return [f"[{target_lang}] {t}" for t in texts]


class TestPptxTranslator(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.test_pptx_path = Path(self.temp_dir.name) / "presentation.pptx"

        prs = Presentation()
        # Slide 1: Title slide
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)

        txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = "Quarterly Business Review"

        p2 = tf.add_paragraph()
        p2.text = "Performance and Outlook"

        prs.save(str(self.test_pptx_path))

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_inspect_file(self):
        info = PptxTranslator.inspect_file(str(self.test_pptx_path))
        self.assertEqual(info["type"], "pptx")
        self.assertEqual(info["slides"], 1)
        self.assertGreater(info["shapes"], 0)
        self.assertGreater(info["word_count"], 3)
        self.assertGreater(info["translatable_segments"], 1)

    def test_translation_flow(self):
        config = AppConfig()
        mock_llm = MockLLMManager()
        cache = TranslationCache(db_path=Path(self.temp_dir.name) / "cache.db")
        translator = PptxTranslator(config, mock_llm, cache=cache)

        out_file = translator.process_file(
            file_path=str(self.test_pptx_path),
            target_lang="Vietnamese"
        )

        out_path = Path(out_file)
        self.assertTrue(out_path.exists())
        self.assertEqual(out_path.name, "presentation_vi.pptx")

        translated_prs = Presentation(str(out_path))
        all_texts = []
        for slide in translated_prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        all_texts.append(p.text)

        joined = " ".join(all_texts)
        self.assertIn("[Vietnamese] Quarterly Business Review", joined)
        self.assertIn("[Vietnamese] Performance and Outlook", joined)


if __name__ == "__main__":
    unittest.main()

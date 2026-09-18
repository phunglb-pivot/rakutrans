"""
PowerPoint (.pptx) parsing and translation engine for RakuTrans AI.
Uses python-pptx to translate slide decks while preserving slide layouts,
shapes, tables, formatting, and speaker notes.
"""

from pathlib import Path
from typing import List, Dict, Set, Callable, Optional, Any
from concurrent.futures import ThreadPoolExecutor, as_completed
import pptx
from pptx import Presentation
from pptx.text.text import _Paragraph

from base_translator import BaseTranslator, safe_save_file
from config import AppConfig, TranslationOptions, get_smart_output_path
from llm_manager import LLMManager
from translation_cache import TranslationCache


def _should_translate(text: Optional[str]) -> bool:
    """Checks if a string has translatable content."""
    if not text:
        return False
    trimmed = text.strip()
    if not trimmed:
        return False
    try:
        float(trimmed.replace(",", ""))
        return False
    except ValueError:
        pass
    return any(c.isalnum() for c in trimmed)


class PptxTranslator(BaseTranslator):
    """
    Handles translation of Microsoft PowerPoint (.pptx) presentations.
    Preserves slide layouts, text boxes, shapes, and tables.
    Features text deduplication, local SQLite cache, and concurrent batching.
    """

    @classmethod
    def _iterate_paragraphs(cls, prs: Presentation):
        """Yields all paragraphs across all slides, shapes, tables, and notes."""
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        yield p
                elif shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            for p in cell.text_frame.paragraphs:
                                yield p
                elif shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
                    try:
                        for sub_shape in shape.shapes:
                            if sub_shape.has_text_frame:
                                for p in sub_shape.text_frame.paragraphs:
                                    yield p
                    except Exception:
                        pass

            # Also check slide notes if present
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                for p in slide.notes_slide.notes_text_frame.paragraphs:
                    yield p

    @classmethod
    def inspect_file(cls, file_path: str) -> Dict[str, Any]:
        """Inspects PowerPoint presentation to count slides, shapes, words, and text segments."""
        path = Path(file_path)
        if not path.exists():
            return {"type": "pptx", "error": "File not found"}

        try:
            prs = Presentation(str(path))
            slide_count = len(prs.slides)
            shape_count = sum(len(s.shapes) for s in prs.slides)
            total_words = 0
            unique_texts: Set[str] = set()

            for p in cls._iterate_paragraphs(prs):
                text = p.text
                if text:
                    total_words += len(text.split())
                    if _should_translate(text):
                        unique_texts.add(text.strip())

            return {
                "type": "pptx",
                "slides": slide_count,
                "shapes": shape_count,
                "word_count": total_words,
                "translatable_segments": len(unique_texts)
            }
        except Exception as e:
            return {"type": "pptx", "error": str(e)}

    def process_file(
        self,
        file_path: str,
        target_lang: Optional[str] = None,
        source_lang: Optional[str] = "auto",
        options: Optional[TranslationOptions] = None,
        progress_callback: Optional[Callable[[int, int, str], None]] = None,
        custom_output_path: Optional[str] = None,
        *args,
        **kwargs
    ) -> str:
        """
        Translates a PowerPoint presentation into the target language.
        Extracts translatable text segments, translates via LLM in concurrent batches,
        and saves output file with layout preservation.
        """
        src = Path(file_path)
        if not src.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        target_lang = target_lang or "Vietnamese"
        options = options or self.config.translation_options

        if progress_callback:
            progress_callback(5, 100, "Reading PowerPoint presentation...")

        prs = Presentation(str(src))

        # Collect paragraphs to translate
        paragraphs_to_process: List[_Paragraph] = []
        for p in self._iterate_paragraphs(prs):
            if _should_translate(p.text):
                paragraphs_to_process.append(p)

        if not paragraphs_to_process:
            out_path = get_smart_output_path(src, target_lang=target_lang, custom_output_path=custom_output_path)
            safe_save_file(lambda p: prs.save(str(p)), out_path)
            if progress_callback:
                progress_callback(100, 100, "Presentation contains no text to translate.")
            return str(out_path)

        unique_texts = list({p.text.strip() for p in paragraphs_to_process if _should_translate(p.text)})
        total_unique = len(unique_texts)

        if progress_callback:
            progress_callback(10, 100, f"Found {total_unique} unique text blocks. Checking cache...")

        translation_map: Dict[str, str] = {}
        uncached_texts: List[str] = []

        for text in unique_texts:
            cached = self.cache.get(text, target_lang)
            if cached:
                translation_map[text] = cached
            else:
                uncached_texts.append(text)

        if uncached_texts:
            batch_size = max(5, getattr(self.config, "batch_size", 20))
            batches = [uncached_texts[i:i + batch_size] for i in range(0, len(uncached_texts), batch_size)]
            total_batches = len(batches)

            completed_batches = 0

            def translate_worker(batch: List[str]) -> List[str]:
                if self.is_cancelled:
                    raise InterruptedError("Translation cancelled by user.")
                return self.llm_manager.translate_batch(
                    texts=batch,
                    target_lang=target_lang,
                    source_lang=source_lang,
                    options=options
                )

            max_workers = min(3, max(1, total_batches))
            with ThreadPoolExecutor(max_workers=max_workers) as executor:
                futures = {executor.submit(translate_worker, b): b for b in batches}
                for future in as_completed(futures):
                    if self.is_cancelled:
                        executor.shutdown(wait=False, cancel_futures=True)
                        raise InterruptedError("Translation cancelled by user.")

                    orig_batch = futures[future]
                    try:
                        translated_batch = future.result()
                        for orig, trans in zip(orig_batch, translated_batch):
                            translation_map[orig] = trans
                            self.cache.set(orig, target_lang, trans)
                    except Exception as e:
                        if self.is_cancelled:
                            raise InterruptedError("Translation cancelled by user.")
                        raise e

                    completed_batches += 1
                    pct = 10 + int(completed_batches / total_batches * 75)
                    if progress_callback:
                        progress_callback(pct, 100, f"Translating batch {completed_batches}/{total_batches}...")

        if self.is_cancelled:
            raise InterruptedError("Translation cancelled by user.")

        if progress_callback:
            progress_callback(90, 100, "Applying translations to PowerPoint presentation...")

        # Apply translations back
        for p in paragraphs_to_process:
            orig_text = p.text.strip()
            translated = translation_map.get(orig_text)
            if not translated:
                continue

            if len(p.runs) <= 1:
                if p.runs:
                    p.runs[0].text = translated
                else:
                    p.text = translated
            else:
                p.runs[0].text = translated
                for r in p.runs[1:]:
                    r.text = ""

        if progress_callback:
            progress_callback(95, 100, "Saving translated presentation...")

        out_path = get_smart_output_path(
            source_path=src,
            target_lang=target_lang,
            target_format=".pptx",
            custom_output_path=custom_output_path
        )
        saved_path = safe_save_file(lambda p: prs.save(str(p)), out_path)

        if progress_callback:
            progress_callback(100, 100, f"Saved: {saved_path.name}")

        return str(saved_path)

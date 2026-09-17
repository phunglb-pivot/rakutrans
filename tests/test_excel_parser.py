"""
Unit tests for ExcelParser (openpyxl) formula and styling preservation.
"""

import os
import tempfile
import unittest
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from config import AppConfig, TranslationOptions
from excel_parser import ExcelTranslator


class MockLLMManager:
    """Mock LLMManager to test translation re-injection deterministically without API calls."""
    def __init__(self):
        self.call_count = 0

    def translate_batch(self, texts, *args, **kwargs):
        self.call_count += 1
        target_lang = kwargs.get("target_lang", "Vietnamese")
        if args:
            if len(args) >= 2 and isinstance(args[1], str):
                target_lang = args[1]
            elif len(args) >= 1 and isinstance(args[0], str):
                target_lang = args[0]
        # Prefix translated text to simulate translation
        return [f"[{target_lang[:2].upper()}] {t}" for t in texts]


class TestExcelTranslator(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.test_excel_path = Path(self.temp_dir.name) / "test_sample.xlsx"

        # Create a sample multi-sheet workbook with formulas, numbers, and styles
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "SummarySheet"

        # Sheet 1: Header with styling
        ws1["A1"] = "Project Specification"
        ws1["A1"].font = Font(name="Arial", size=14, bold=True, color="FF0000")
        ws1["A1"].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

        # Plain numbers (must NOT be translated)
        ws1["A2"] = 100
        ws1["A3"] = 250.75

        # Formula (must NOT be translated, MUST be preserved exactly as string "=SUM(A2:A3)")
        ws1["A4"] = "=SUM(A2:A3)"

        # Translatable strings
        ws1["B2"] = "Database Migration"
        ws1["B3"] = "Authentication Service"

        # Empty cell
        ws1["C2"] = None
        ws1["C3"] = "   "

        # Sheet 2
        ws2 = wb.create_sheet(title="DetailsSheet")
        ws2["A1"] = "API Endpoint Architecture"
        ws2["B1"] = "=COUNTIF(A1:A10, \">0\")"
        ws2["A2"] = 42

        wb.save(self.test_excel_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_cell_filtering_rules(self):
        config = AppConfig()
        mock_llm = MockLLMManager()
        translator = ExcelTranslator(config, mock_llm)

        # Numbers
        self.assertFalse(translator.should_translate_cell(100))
        self.assertFalse(translator.should_translate_cell(45.67))
        self.assertFalse(translator.should_translate_cell("12345"))
        self.assertFalse(translator.should_translate_cell("1,234.56"))

        # Empty / Whitespace
        self.assertFalse(translator.should_translate_cell(None))
        self.assertFalse(translator.should_translate_cell(""))
        self.assertFalse(translator.should_translate_cell("   "))

        # Formulas
        self.assertFalse(translator.should_translate_cell("=SUM(A1:A10)"))
        self.assertFalse(translator.should_translate_cell("=IF(B2>0, 1, 0)"))

        # Valid text
        self.assertTrue(translator.should_translate_cell("User Authentication"))
        self.assertTrue(translator.should_translate_cell("仕様書の確認"))

    def test_formula_and_style_preservation(self):
        config = AppConfig(batch_size=50)
        mock_llm = MockLLMManager()
        translator = ExcelTranslator(config, mock_llm)

        output_path = translator.process_file(
            file_path=str(self.test_excel_path),
            source_lang="Japanese",
            target_lang="Vietnamese"
        )

        self.assertTrue(Path(output_path).exists())

        # Verify translated workbook
        result_wb = openpyxl.load_workbook(output_path, data_only=False)

        # Check sheets preserved
        self.assertEqual(result_wb.sheetnames, ["SummarySheet", "DetailsSheet"])

        ws1 = result_wb["SummarySheet"]

        # 1. Check style preservation on A1
        self.assertEqual(ws1["A1"].value, "[VI] Project Specification")
        self.assertTrue(ws1["A1"].font.bold)
        self.assertEqual(ws1["A1"].font.color.rgb, "00FF0000")
        self.assertEqual(ws1["A1"].fill.start_color.rgb, "00FFFF00")

        # 2. Check numbers untouched
        self.assertEqual(ws1["A2"].value, 100)
        self.assertEqual(ws1["A3"].value, 250.75)

        # 3. Check formula STRICTLY preserved
        self.assertEqual(ws1["A4"].value, "=SUM(A2:A3)")

        # 4. Check strings translated
        self.assertEqual(ws1["B2"].value, "[VI] Database Migration")
        self.assertEqual(ws1["B3"].value, "[VI] Authentication Service")

        # 5. Check empty cells remain empty
        self.assertIsNone(ws1["C2"].value)
        self.assertEqual(ws1["C3"].value, "   ")

        # 6. Check Sheet 2
        ws2 = result_wb["DetailsSheet"]
        self.assertEqual(ws2["A1"].value, "[VI] API Endpoint Architecture")
        self.assertEqual(ws2["B1"].value, "=COUNTIF(A1:A10, \">0\")")
        self.assertEqual(ws2["A2"].value, 42)

    def test_in_file_deduplication(self):
        # Create workbook where same text appears multiple times
        dedup_path = Path(self.temp_dir.name) / "dedup_sample.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DeduplicationTest"

        # 10 cells with identical text "Status: Done"
        for i in range(1, 11):
            ws[f"A{i}"] = "Status: Done"
        # 5 cells with identical text "Priority: High"
        for i in range(1, 6):
            ws[f"B{i}"] = "Priority: High"

        wb.save(dedup_path)

        class CountingMockLLM(MockLLMManager):
            def __init__(self):
                super().__init__()
                self.translated_texts_history = []

            def translate_batch(self, texts, *args, **kwargs):
                self.translated_texts_history.extend(texts)
                return super().translate_batch(texts, *args, **kwargs)

        from translation_cache import TranslationCache
        temp_cache = TranslationCache(db_path=Path(self.temp_dir.name) / "test_dedup_cache.db")
        mock_llm = CountingMockLLM()
        config = AppConfig(batch_size=50)
        translator = ExcelTranslator(config, mock_llm, cache=temp_cache)

        out = translator.process_file(
            file_path=str(dedup_path),
            target_lang="Vietnamese"
        )

        # 15 cells total, but only 2 unique texts
        self.assertEqual(len(mock_llm.translated_texts_history), 2)
        self.assertIn("Status: Done", mock_llm.translated_texts_history)
        self.assertIn("Priority: High", mock_llm.translated_texts_history)

        # Verify all 15 cells in output workbook have correct translation
        res_wb = openpyxl.load_workbook(out)
        res_ws = res_wb.active
        for i in range(1, 11):
            self.assertEqual(res_ws[f"A{i}"].value, "[VI] Status: Done")
        for i in range(1, 6):
            self.assertEqual(res_ws[f"B{i}"].value, "[VI] Priority: High")

    def test_inspect_file(self):
        info = ExcelTranslator.inspect_file(str(self.test_excel_path))
        self.assertEqual(info["type"], "excel")
        self.assertEqual(info["sheet_count"], 2)
        self.assertIn("SummarySheet", info["sheet_names"])
        self.assertIn("DetailsSheet", info["sheet_names"])
        self.assertGreater(info["translatable_cells"], 0)
        self.assertGreater(info["formula_count"], 0)

    def test_bilingual_mode(self):
        mock_llm = MockLLMManager()
        config = AppConfig(batch_size=50)
        translator = ExcelTranslator(config, mock_llm)
        options = TranslationOptions(bilingual_mode=True)

        out = translator.process_file(
            file_path=str(self.test_excel_path),
            target_lang="Vietnamese",
            options=options
        )

        res_wb = openpyxl.load_workbook(out)
        # Should contain original sheets + duplicated archived sheets
        self.assertIn("SummarySheet", res_wb.sheetnames)
        self.assertIn("DetailsSheet", res_wb.sheetnames)
        self.assertTrue(any("SummarySheet_Orig" in name for name in res_wb.sheetnames))
        self.assertTrue(any("DetailsSheet_Orig" in name for name in res_wb.sheetnames))

    def test_safe_save_workbook(self):
        from excel_parser import safe_save_workbook
        wb = openpyxl.Workbook()
        target = Path(self.temp_dir.name) / "locked_file.xlsx"
        # First save succeeds
        out1 = safe_save_workbook(wb, target)
        self.assertEqual(out1, target)

        # Mock permission error on locked target
        orig_save = wb.save
        def mock_save(path):
            if str(path) == str(target):
                raise PermissionError("File in use by Microsoft Excel")
            return orig_save(path)

        wb.save = mock_save
        out2 = safe_save_workbook(wb, target)
        self.assertIn("(1)", str(out2))
        self.assertTrue(out2.exists())


if __name__ == "__main__":
    unittest.main()

